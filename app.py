import streamlit as st
from google import genai
from google.genai import types
import json
import io
import zipfile

# Bibliotecas para extração de dados do Microsoft Office
import docx2txt
import openpyxl
from pptx import Presentation

st.set_page_config(page_title="Innpulse Fórum 2026 - MVP", page_icon="📁", layout="wide")
st.title("Motor de Classificação - Innpulse Fórum 2026 📁")
st.write("Submeta múltiplos arquivos para organização automatizada baseada na inteligência de negócio do ecossistema.")

def extrair_texto_office(file, extensao):
    texto = ""
    try:
        if extensao == "docx":
            texto = docx2txt.process(file)
        elif extensao == "xlsx":
            wb = openpyxl.load_workbook(file, data_only=True)
            for sheet in wb.worksheets:
                texto += f"\nAba: {sheet.title}\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) for cell in row if cell is not None])
                    if row_text:
                        texto += row_text + "\n"
        elif extensao == "pptx":
            prs = Presentation(file)
            for i, slide in enumerate(prs.slides):
                texto += f"\nSlide {i+1}:\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texto += shape.text + "\n"
    except Exception as e:
        texto = f"Erro na extração do arquivo Office: {str(e)}"
    return texto

uploaded_files = st.file_uploader(
    "Arraste ou selecione os arquivos do Innpulse Fórum", 
    type=["pdf", "png", "jpg", "jpeg", "docx", "xlsx", "pptx", "txt"],
    accept_multiple_files=True
)

if uploaded_files:
    st.write(f"📂 **{len(uploaded_files)} arquivos carregados no buffer.**")
    
    if st.button("🚀 Executar Ingestão Inteligente"):
        client = genai.Client(api_key=st.secrets.get("GEMINI_API_KEY", "SUA_CHAVE_AQUI"))
        zip_buffer = io.BytesIO()
        arquivos_processados = []
        erros = []
        
        progresso_barra = st.progress(0)
        status_texto = st.empty()
        
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            for index, file in enumerate(uploaded_files):
                nome_original = file.name
                extensao = nome_original.split(".")[-1].lower()
                
                progresso_barra.progress((index + 1) / len(uploaded_files))
                status_texto.text(f"Analisando documento ({index + 1}/{len(uploaded_files)}): {nome_original}")
                
                # SYSTEM PROMPT EXTRAÍDO FIELMENTE DO SEU DOCUMENTO DOCX
                prompt_sistema = (
                    "Você é o motor de classificação de arquivos do MVP de Gestão de Conhecimento do Innpulse Fórum 2026. "
                    "Sua tarefa é ler o conteúdo de um arquivo que foi submetido, entender seu contexto e retornar estritamente "
                    "um objeto JSON com a classificação correta, seguindo as regras de negócio fornecidas.\n\n"
                    "--- REGRAS DE TAXONOMIA (PASTAS) ---\n"
                    "A pasta raiz sempre será: '01_Innpulse_Forum_2026'\n"
                    "Você deve escolher OBRIGATORIAMENTE uma das seguintes subpastas:\n"
                    "1. '01_Premio_Innpulse' (Se o texto focar em premiação, jurados, categorias e critérios do prêmio)\n"
                    "2. '02_Batalha_Startups_GITR' (Se o texto focar na competição Get in the ring, pitches, ringue, chaves e duelos)\n"
                    "3. '03_Rodada_Negocios' (Se o texto focar em matchmaking, reuniões entre corporações/investidores e startups, agendas)\n\n"
                    "--- REGRAS DE TIPO DE DOCUMENTO ---\n"
                    "Você deve identificar o tipo de documento baseado estritamente nestas opções:\n"
                    "- 'Formulario_Inscricao'\n- 'Comunicacao'\n- 'Edital'\n- 'Roteiro'\n- 'Gestao_Finalistas'\n- 'Relatorio'\n\n"
                    "--- REGRAS DE NOMENCLATURA (PADRÃO) ---\n"
                    "O nome sugerido do arquivo deve seguir RIGOROSAMENTE o formato:\n"
                    "YYYY_MM_DD_SubpastaSemNumero_TipoDeDocumento_Resumo_V1.[extensão_original]\n"
                    "Notas importantes sobre a nomenclatura:\n"
                    "- Procure no texto a data de criação ou a data do evento citada. Se não houver nenhuma data explícita no texto, utilize estritamente a data de hoje: 2026_05_21.\n"
                    "- No campo 'SubpastaSemNumero', remova os dígitos iniciais e use apenas o termo descritivo correspondente: PremioInnpulse, BatalhaGITR, ou RodadaNegocios.\n"
                    "- O campo 'Resumo' deve ser curto (1 a 3 palavras), utilizando formato CamelCase, sem espaços.\n\n"
                    "--- FORMATO DE SAÍDA OBRIGATÓRIO (JSON) ---\n"
                    "Sua resposta deve seguir exatamente este mapeamento de chaves:\n"
                    "{\n"
                    "  \"pasta_raiz\": \"01_Innpulse_Forum_2026\",\n"
                    "  \"subpasta\": \"NOME_DA_SUBPASTA_ESCOLHIDA\",\n"
                    "  \"tipo_documento\": \"TIPO_IDENTIFICADO\",\n"
                    "  \"nome_sugerido_arquivo\": \"NOME_PADRONIZADO.ext\",\n"
                    "  \"palavras_chave_banco_dados\": [\"termo1\", \"termo2\", \"termo3\", \"termo4\"],\n"
                    "  \"resumo_conteudo\": \"Frase curta resumindo os dados estruturados encontrados.\"\n"
                    "}"
                )
                
                # Encapsulamento multimodal ou textual conforme a extensão detectada
                if extensao in ["pdf", "png", "jpg", "jpeg"]:
                    bytes_data = file.getvalue()
                    mime_type = f"application/{extensao}" if extensao == "pdf" else f"image/{extensao}"
                    if extensao == "jpg": mime_type = "image/jpeg"
                    conteudo_para_ia = [types.Part.from_bytes(data=bytes_data, mime_type=mime_type), prompt_sistema]
                else:
                    texto_extraido = file.read().decode("utf-8") if extensao == "txt" else extrair_texto_office(file, extensao)
                    conteudo_para_ia = [f"Texto bruto para análise do motor:\n{texto_extraido}\n\n", prompt_sistema]
                
                try:
                    # Executa a chamada forçando a saída tipada em JSON para evitar quebras
                    response = client.models.generate_content(
                        model='gemini-2.5-flash',
                        contents=conteudo_para_ia,
                        config=types.GenerateContentConfig(
                            response_mime_type="application/json"
                        )
                    )
                    
                    resultado = json.loads(response.text.strip())
                    
                    # Consome as chaves estruturadas mapeadas no documento original
                    raiz = resultado['pasta_raiz'].strip()
                    subpasta = resultado['subpasta'].strip()
                    nome_final = resultado['nome_sugerido_arquivo'].strip()
                    
                    # Gera a árvore virtual dentro do pacote binário compactado
                    caminho_no_zip = f"{raiz}/{subpasta}/{nome_final}"
                    zip_file.writestr(caminho_no_zip, file.getvalue())
                    
                    arquivos_processados.append({
                        "original": nome_original,
                        "novo": nome_final,
                        "pasta": f"{raiz}/{subpasta}",
                        "tipo": resultado['tipo_documento'],
                        "tags": ", ".join(resultado['palavras_chave_banco_dados']),
                        "resumo": resultado['resumo_conteudo']
                    })
                    
                except Exception as e:
                    # Aloca arquivos com erro em uma subpasta isolada dentro do diretório raiz do projeto
                    zip_file.writestr(f"01_Innpulse_Forum_2026/ERROS_PROCESSAMENTO/{nome_original}", file.getvalue())
                    erros.append({"arquivo": nome_original, "erro": str(e)})
        
        status_texto.success("✨ Processamento do lote concluído com sucesso!")
        progresso_barra.empty()
        
        st.write("---")
        st.subheader("📦 Download do Repositório Estruturado")
        st.download_button(
            label="📥 BAIXAR ESTRUTURA DE PASTAS COMPLETA (.ZIP)",
            data=zip_buffer.getvalue(),
            file_name="01_Innpulse_Forum_2026.zip",
            mime="application/zip",
            use_container_width=True
        )
        
        # Consolidação de Metadados voltada para a camada de Consulta Semântica descrita no MVP
        st.write("---")
        st.subheader("📊 Relatório de Metadados Extraídos (Pronto para Banco de Dados)")
        for item in arquivos_processados:
            with st.expander(f"📄 {item['original']} ➔ {item['novo']}"):
                st.markdown(f"**📍 Armazenamento Direcionado:** `{item['pasta']}`")
                st.markdown(f"**🏷️ Tipo de Documentação:** `{item['tipo']}`")
                st.markdown(f"**🔑 Indexadores de Busca Semântica:** *{item['tags']}*")
                st.markdown(f"**📝 Resumo do Conteúdo Encontrado:** {item['resumo']}")
                
        if erros:
            st.warning("⚠️ Arquivos direcionados para a pasta de contingência (/ERROS_PROCESSAMENTO/):")
            for err in erros:
                st.write(f"❌ Documento: {err['arquivo']} | Falha detectada: {err['erro']}")
